from django.shortcuts import get_object_or_404
from django.views import View
from django.http import JsonResponse
from django.db import transaction
from django.db.models import Q, OuterRef, Subquery, Max
from api.models import UserMaster, FactoryItemIn, DailyWorkOrder, DailyWorker
from api.lib import Pagenation, get_excep_msg
from django.utils import timezone
from api.factory.factory_lib import get_itemin_code, get_user_info, get_workorder_code
import json
from datetime import timedelta
import calendar
from django.utils.dateparse import parse_date
from django.db.models import Sum
from django.db.models.functions import TruncMonth, TruncDay
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse
from io import BytesIO

"""
class MonthlyClaimReport_ExcelExport(View):
    def get(self, request, *args, **kwargs):
        fr_date = request.GET.get('fr_date', '')
        to_date = request.GET.get('to_date', '')
        sch_subsidiary = request.GET.get('sch_subsidiary', '')

        qs = DailyWorker.objects.filter(company=request.user.company).select_related(
            "work_order__subsidiary"
        )

        if fr_date:
            qs = qs.filter(work_order__date__gte=fr_date)
        if to_date:
            qs = qs.filter(work_order__date__lte=to_date)
        if sch_subsidiary:
            qs = qs.filter(work_order__subsidiary__id=sch_subsidiary)

        # subsidiary / customer 제거 → user 단위로만 취합
        qs = (
            qs.select_related("user")
            .annotate(month=TruncMonth("work_order__date"))
            .values(
                "month",
                "user__id",     # 그룹 기준
                "user__name",   # 표시용
                "user__rrn",
            )
            .annotate(
                total_pre_result_price=Sum("pre_result_price"),
            )
            .order_by(
                "month",
                "user__id",  # 사용자 기준 정렬
            )
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'sheet1'

        # 헤더 정의
        headers = ['이름', '주민번호', '지급총액']
        ws.append(headers)

        for row in qs:
            ws.append([
                row["user__name"],  # 이름
                row["user__rrn"],   # 주민번호
                row["total_pre_result_price"] or 0,  # 지급총액
            ])

        # 헤더 스타일
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 금액 부분 셀 서식 (3번째 열)
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
            for cell in row:
                cell.number_format = '#,##0'  # 천 단위 콤마, 소수점 없음

        # 필터 적용
        ws.auto_filter.ref = ws.dimensions

        # 열 너비 자동 조정 (한글 보정 포함)
        for column_cells in ws.columns:
            max_length = 0
            col = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    cell_value = str(cell.value) if cell.value else ''
                    display_length = sum(2 if ord(ch) > 127 else 1 for ch in cell_value)
                    max_length = max(max_length, display_length)
                except:
                    pass
            ws.column_dimensions[col].width = max_length + 3

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        filename = f"monthly_claim_report_{fr_date or 'start'}_{to_date or 'end'}.xlsx"
        response['Content-Disposition'] = f'attachment; filename={filename}'
        return response
"""

# 초기코드 (법인별, 센터별 근무자의 공제액, 실제급액 정보)
class MonthlyClaimReport_ExcelExport(View):
     def get(self, request, *args, **kwargs):
         fr_date = request.GET.get('fr_date', '')
         to_date = request.GET.get('to_date', '')
         sch_subsidiary = request.GET.get('sch_subsidiary', '')
         sch_customer = request.GET.get('sch_customer', '')

         qs = DailyWorker.objects.filter(company=request.user.company)

         if fr_date:
             qs = qs.filter(work_order__date__gte=fr_date)
         if to_date:
             qs = qs.filter(work_order__date__lte=to_date)
         if sch_subsidiary:
            qs = qs.filter(work_order__subsidiary__id=sch_subsidiary)
         if sch_customer:
             qs = qs.filter(work_order__customer_id=sch_customer)

         qs = (
             qs.select_related("work_order__subsidiary", "work_order__customer", "user")
             .annotate(month=TruncMonth("work_order__date"))
             .values(
                 "month",
                 "work_order__subsidiary__name",
                 "work_order__customer__name",
                 "user__id",  # 그룹 기준: 사용자 ID 포함
                 "user__name",  # 표시용
             )
             .annotate(
                 total_tax=Sum("tax"),
                 total_result_price=Sum("result_price"),
             )
             .order_by(
                 "month",
                 "work_order__subsidiary__name",
                 "work_order__customer__name",
                 "user__id",  # 정렬도 id 기준
             )
         )

         wb = openpyxl.Workbook()
         ws = wb.active
         ws.title = 'sheet1'

         # 헤더 정의
         headers = ['순서', '연월', '관리법인', '센터', '근무자 이름', '총 공제액', '총 실지급액']
         ws.append(headers)

         for idx, row in enumerate(qs, start=1):
             ws.append([
                 idx,
                 row["month"].strftime("%Y-%m"),
                 row["work_order__subsidiary__name"] or '',
                 row["work_order__customer__name"] or '',
                 row["user__name"],  # 이름만 출력
                 row["total_tax"] or 0,
                 row["total_result_price"] or 0,
             ])

         # 헤더 스타일
         header_font = Font(bold=True, color="FFFFFF")
         header_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
         for cell in ws[1]:
             cell.font = header_font
             cell.fill = header_fill
             cell.alignment = Alignment(horizontal="center", vertical="center")

         # 금액 부분 셀 서식 (F, G 컬럼)
         for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):  # 6=F(공제액), 7=G(실지급액)
             for cell in row:
                 cell.number_format = '#,##0'  # 천 단위 콤마, 소수점 없음

         # 필터 적용
         ws.auto_filter.ref = ws.dimensions

         # 열 너비 자동 조정 (한글 보정 포함)
         for column_cells in ws.columns:
             max_length = 0
             col = column_cells[0].column_letter
             for cell in column_cells:
                 try:
                     cell_value = str(cell.value) if cell.value else ''
                     display_length = sum(2 if ord(ch) > 127 else 1 for ch in cell_value)
                     max_length = max(max_length, display_length)
                 except:
                     pass
             ws.column_dimensions[col].width = max_length + 3

         buffer = BytesIO()
         wb.save(buffer)
         buffer.seek(0)

         response = HttpResponse(buffer.read(),
                                 content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
         filename = f"monthly_claim_report_{fr_date or 'start'}_{to_date or 'end'}.xlsx"
         response['Content-Disposition'] = f'attachment; filename={filename}'
         return response


# 초기코드 (법인별, 센터별 청구금액 정보)
class MonthlyTotalReport_ExcelExport(View):
     def get(self, request, *args, **kwargs):
         fr_date = request.GET.get('fr_date', '')
         to_date = request.GET.get('to_date', '')
         sch_subsidiary = request.GET.get('sch_subsidiary', '')
         sch_customer = request.GET.get('sch_customer', '')

         qs = DailyWorker.objects.filter(company=request.user.company)

         if fr_date:
             qs = qs.filter(work_order__date__gte=fr_date)
         if to_date:
             qs = qs.filter(work_order__date__lte=to_date)
         if sch_subsidiary:
            qs = qs.filter(work_order__subsidiary__id=sch_subsidiary)
         if sch_customer:
             qs = qs.filter(work_order__customer_id=sch_customer)
         qs = (
             qs.select_related("work_order__subsidiary", "work_order__customer", "user")
             .annotate(truncday=TruncDay("work_order__date"))
             .values(
                 "truncday",
                 "work_order__subsidiary__name",
                 "work_order__customer__name",
                 "user__id",  # 그룹 기준: 사용자 ID 포함
                 "user__name",  # 표시용
             )
             .annotate(
                 total_tax=Sum("tax"),
                 total_total=Sum("total"),
             )
             .order_by(
                 "work_order__subsidiary__name",
                 "work_order__customer__name",
                 "truncday",
                 "user__id",  # 정렬도 id 기준
             )
         )

         wb = openpyxl.Workbook()
         ws = wb.active
         ws.title = 'sheet1'

         # 헤더 정의
         headers = ['순서', '관리법인', '센터', '일자', '근무자 이름', '총 청구금액']
         ws.append(headers)

         for idx, row in enumerate(qs, start=1):
             ws.append([
                 idx,
                 row["work_order__subsidiary__name"] or '',
                 row["work_order__customer__name"] or '',
                 row["truncday"].strftime("%Y-%m-%d"),
                 row["user__name"],  # 이름만 출력
                 row["total_total"] or 0,
             ])

         # 헤더 스타일
         header_font = Font(bold=True, color="FFFFFF")
         header_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")
         for cell in ws[1]:
             cell.font = header_font
             cell.fill = header_fill
             cell.alignment = Alignment(horizontal="center", vertical="center")

         # 금액 부분 셀 서식 (F 컬럼)
         for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):  # 6=F(청구금액)
             for cell in row:
                 cell.number_format = '#,##0'  # 천 단위 콤마, 소수점 없음

         # 필터 적용
         ws.auto_filter.ref = ws.dimensions

         # 열 너비 자동 조정 (한글 보정 포함)
         for column_cells in ws.columns:
             max_length = 0
             col = column_cells[0].column_letter
             for cell in column_cells:
                 try:
                     cell_value = str(cell.value) if cell.value else ''
                     display_length = sum(2 if ord(ch) > 127 else 1 for ch in cell_value)
                     max_length = max(max_length, display_length)
                 except:
                     pass
             ws.column_dimensions[col].width = max_length + 3

         buffer = BytesIO()
         wb.save(buffer)
         buffer.seek(0)

         response = HttpResponse(buffer.read(),
                                 content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
         filename = f"monthly_total_report_{fr_date or 'start'}_{to_date or 'end'}.xlsx"
         response['Content-Disposition'] = f'attachment; filename={filename}'
         return response


from urllib.parse import quote
import re
class MonthlyClaimBank_ExcelExport(View):
    def get(self, request, *args, **kwargs):
        fr_date = request.GET.get('fr_date', '')
        to_date = request.GET.get('to_date', '')
        sch_subsidiary = request.GET.get('sch_subsidiary', '')
        sch_customer = request.GET.get('sch_customer', '')
        qs = DailyWorker.objects.filter(company=request.user.company).select_related(
            "work_order__subsidiary", "user__bank"
        )

        if sch_subsidiary:
            qs = qs.filter(work_order__subsidiary__id=sch_subsidiary)
        if sch_customer:
            qs = qs.filter(work_order__customer_id=sch_customer)

        if fr_date:
            qs = qs.filter(work_order__date__gte=fr_date)
        if to_date:
            qs = qs.filter(work_order__date__lte=to_date)
        if sch_subsidiary:
            qs = qs.filter(work_order__subsidiary__id=sch_subsidiary)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'sheet1'

        for obj in qs:
            row = [
                obj.user.bank.name if obj.user and obj.user.bank else '',  # 은행명
                obj.user.account_code if obj.user else '',  # 계좌번호
                obj.result_price or 0,  # 실지급액
                obj.user.name if obj.user else '',  # 근무자 이름
                f"(주) {obj.work_order.subsidiary.name}" if obj.work_order and obj.work_order.subsidiary else '',  # 관리법인
            ]

            # 관리법인명에 '이음'이 포함되어 있다면 컬럼 2개 추가
            if obj.work_order and obj.work_order.subsidiary and "이음" in obj.work_order.subsidiary.name:
                row.extend(["일급", "일급"])

            ws.append(row)

        # 실지급액 (3번째 컬럼)만 숫자 서식 적용
        for row in ws.iter_rows(min_row=1, min_col=3, max_col=3):
            for cell in row:
                cell.number_format = '#,##0'

        # 열 너비 자동 조정
        for column_cells in ws.columns:
            max_length = 0
            col = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    cell_value = str(cell.value) if cell.value else ''
                    display_length = sum(2 if ord(ch) > 127 else 1 for ch in cell_value)
                    max_length = max(max_length, display_length)
                except:
                    pass
            ws.column_dimensions[col].width = max_length + 3

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

        # 안전한 subsidiary_name 만들기
        subsidiary_name = ""
        if qs.exists():
            first_obj = qs.first()
            if first_obj and first_obj.work_order and first_obj.work_order.subsidiary:
                raw_name = first_obj.work_order.subsidiary.name
                # 파일명에 쓸 수 없는 문자 제거/치환
                safe_name = re.sub(r'[\\/*?:"<>|]', "_", raw_name).replace(" ", "_")
                subsidiary_name = safe_name

        filename = f"monthly_claim_bank_{fr_date or 'start'}_{to_date or 'end'}"
        if subsidiary_name:
            filename += f"_{subsidiary_name}"
        filename += ".xlsx"

        # Content-Disposition 헤더 (UTF-8 대응)
        response['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"

        # filename = f"monthly_claim_bank_{fr_date or 'start'}_{to_date or 'end'}.xlsx"
        # response['Content-Disposition'] = f'attachment; filename={filename}'
        return response
