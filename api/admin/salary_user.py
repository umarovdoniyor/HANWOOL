from django.views import View
from django.http import JsonResponse
from api.models import UserMaster, UserBaseSalary, SalaryMaster, SalaryItem
from api.lib import Pagenation, get_excep_msg
from django.db.models import Q, F
from django.db import transaction, IntegrityError
from django.utils import timezone
from django.shortcuts import get_object_or_404
import json
from django.utils.decorators import method_decorator
from django.contrib.auth.decorators import login_required
from django.views.decorators.csrf import csrf_exempt
import openpyxl
from openpyxl.styles import Font, PatternFill
from django.http import HttpResponse
from io import BytesIO
from api.basic_data.notification import noti_create_fn


# 직원 기본급여 설정 함수 -------------------------------------------------------------------------------------------------

class UserSalary_List(View):
    def get(self, request, *args, **kwargs):
        is_staff = request.GET.get("is_staff", "").strip()

        if not request.user.is_master:  # 관리자외 접근금지
            # qs = None
            return JsonResponse({'error': '접근 권한이 없습니다.'}, status=403)

        else:
            if is_staff == "true":
                qs = (
                    UserMaster.objects.filter(company=request.user.company, is_delete=0, is_superuser=0, is_observer=0, work_type="관리직")
                    .order_by('team_id', F('seq_order').asc(nulls_last=True), 'join_date'))
            else:
                qs = (UserMaster.objects.filter(company=request.user.company, is_staff=1, is_delete=0, is_superuser=0, work_type="관리직")
                      .order_by('team_id', F('seq_order').asc(nulls_last=True), 'join_date'))

        # all_sch 검색
        all_sch = request.GET.get("all_sch", "").strip()
        if all_sch:
            filters = (
                    Q(employee_code__icontains=all_sch) |
                    Q(name__icontains=all_sch) |
                    Q(user_id__icontains=all_sch) |
                    Q(team__name__icontains=all_sch) |
                    Q(job_title__name__icontains=all_sch) |
                    Q(job_level__name__icontains=all_sch) |
                    Q(hire_type__icontains=all_sch) |
                    Q(phone__icontains=all_sch) |
                    Q(remark__icontains=all_sch) |
                    Q(company__name__icontains=all_sch)
            )

            if all_sch.isdigit() and len(all_sch) == 4:  # 입력값이 숫자이고 4자리이면 연도 검색
                filters |= Q(join_date__year=int(all_sch))  # join_date의 연도가 입력값과 일치하는 경우

            qs = qs.filter(filters)

        # Pagination
        _page = int(request.GET.get('page', 1)) if request.GET.get('page', '1').isdigit() else 1
        _size = int(request.GET.get('page_size', 10)) if request.GET.get('page_size', '10').isdigit() else 10
        qs_ps = Pagenation(qs, _size, _page)

        pre_page = _page - 1
        url_pre = f"/?page_size={_size}&page={pre_page}" if pre_page >= 1 else None
        next_page = _page + 1
        url_next = f"/?page_size={_size}&page={next_page}" if next_page <= qs_ps.paginator.num_pages else None

        results = [get_obj(row) for row in qs_ps]
        context = {
            'count': qs_ps.paginator.count,
            'previous': url_pre,
            'next': url_next,
            'results': results,
        }
        return JsonResponse(context, safe=False)


class UserSalary_Update(View):
    @transaction.atomic
    def post(self, request, *args, **kwargs):
        try:
            user_id = request.POST.get('user', '')
            user = get_object_or_404(UserMaster, id=user_id)

            # 숫자형 필드 파싱
            def get_int(key):
                val = request.POST.get(key)
                return int(val) if val and val.isdigit() else 0

            base_salary = get_int('base_salary')
            extend_salary = get_int('extend_salary')
            night_salary = get_int('night_salary')
            holiday_salary = get_int('holiday_salary')
            meal_salary = get_int('meal_salary')
            drive_salary = get_int('drive_salary')
            bonus = get_int('bonus')
            incentive = get_int('incentive')
            etc_p1v = get_int('etc_p1v')
            etc_p2v = get_int('etc_p2v')
            etc_p3v = get_int('etc_p3v')
            etc_p4v = get_int('etc_p4v')

            pension = get_int('pension')
            health = get_int('health')
            care = get_int('care')
            unemployment = get_int('unemployment')
            income_tax = get_int('income_tax')
            local_tax = get_int('local_tax')
            etc_m1v = get_int('etc_m1v')
            etc_m2v = get_int('etc_m2v')
            etc_m3v = get_int('etc_m3v')
            etc_m4v = get_int('etc_m4v')
            etc_m5v = get_int('etc_m5v')
            etc_m6v = get_int('etc_m6v')

            # 계산
            total_salary = (
                    base_salary + extend_salary + night_salary + holiday_salary + meal_salary + drive_salary + bonus +
                    incentive + etc_p1v + etc_p2v + etc_p3v + etc_p4v
            )

            total_deduction = (
                    pension + health + care + unemployment + income_tax + local_tax +
                    etc_m1v + etc_m2v + etc_m3v + etc_m4v + etc_m5v + etc_m6v
            )

            real_salary = total_salary - total_deduction

            obj, created = UserBaseSalary.objects.update_or_create(
                user=user,
                defaults={
                    'annual_salary': request.POST.get('annual_salary') or None,

                    'base_salary': base_salary,
                    'extend_salary': extend_salary,
                    'night_salary': night_salary,
                    'holiday_salary': holiday_salary,
                    'meal_salary': meal_salary,
                    'drive_salary': drive_salary,
                    'bonus': bonus,
                    'incentive': incentive,
                    'etc_p1v': etc_p1v,
                    'etc_p2v': etc_p2v,
                    'etc_p3v': etc_p3v,
                    'etc_p4v': etc_p4v,

                    'pension': pension,
                    'health': health,
                    'care': care,
                    'unemployment': unemployment,
                    'income_tax': income_tax,
                    'local_tax': local_tax,
                    'etc_m1v': etc_m1v,
                    'etc_m2v': etc_m2v,
                    'etc_m3v': etc_m3v,
                    'etc_m4v': etc_m4v,
                    'etc_m5v': etc_m5v,
                    'etc_m6v': etc_m6v,

                    'total_salary': total_salary,
                    'total_deduction': total_deduction,
                    'real_salary': real_salary,

                    'created_by': request.user,
                    'updated_by': request.user,
                    'company': request.user.company,
                }
            )

            return JsonResponse({'success': True, 'id': obj.id})

        except Exception as e:
            transaction.set_rollback(True)
            msg = get_excep_msg(e)
            return JsonResponse({'error': True, 'message': msg}, status=500)


# 월급 정산 대상목록 호출함수
class MonthlySalary_List(View):
    def post(self, request, *args, **kwargs):
        ids_raw = request.POST.get("ids", "")
        ids = ids_raw.split(",") if ids_raw else []

        qs = UserBaseSalary.objects.filter(user_id__in=ids, company=request.user.company)
        results = [get_UBS_obj(row) for row in qs]
        return JsonResponse({'results': results}, safe=False)


def get_obj(obj):
    try:
        base_salary = obj.base_salary.latest('created_at')  # related_name='base_salary'
    except UserBaseSalary.DoesNotExist:
        base_salary = None

    return {
        'id': obj.id,
        'user_id': obj.user_id if obj.user_id is not None else '',  # 사용자 ID
        'password': obj.password if obj.user_id is not None else '',  # 비밀번호
        'name': obj.name if obj.name is not None else '',  # 사용자 이름
        'employee_code': obj.employee_code if obj.employee_code is not None else '',  # 사용자 사번
        'email': obj.email if obj.email is not None else '',
        'phone': obj.phone if obj.phone is not None else '',
        'address': obj.address if obj.address is not None else '',
        'hire_type': obj.hire_type if obj.hire_type is not None else '',
        'join_date': obj.join_date if obj.join_date is not None else '',
        'retire_date': obj.retire_date if obj.retire_date is not None else '',
        'remark': obj.remark if obj.remark is not None else '',
        'is_active': obj.is_active if obj.is_active is not None else '',
        'is_delete': obj.is_delete if obj.is_delete is not None else '',
        'is_staff': obj.is_staff if obj.is_staff is not None else '',
        'is_master': obj.is_master if obj.is_master is not None else '',
        'profile_image': obj.profile_image.url if obj.profile_image and obj.profile_image.name else '',
        'profile_doc': obj.profile_doc.name if obj.profile_doc is not None else '',
        'seq_order': obj.seq_order if obj.seq_order is not None else '',
        'last_login': obj.last_login.strftime('%Y-%m-%d %H:%M') if obj.last_login else '',

        'team_id': safe_get(obj, 'team.id'),
        'team_name': safe_get(obj, 'team.name'),
        'job_title_id': safe_get(obj, 'job_title.id'),
        'job_title_name': safe_get(obj, 'job_title.name'),
        'job_level_id': safe_get(obj, 'job_level.id'),
        'job_level_name': safe_get(obj, 'job_level.name'),

        'created_by_id': safe_get(obj, 'created_by.id'),
        'created_by_name': safe_get(obj, 'created_by.name'),
        'created_at': obj.created_at if obj.created_at is not None else '',
        'company_id': safe_get(obj, 'company.id'),
        'company_name': safe_get(obj, 'company.company_info.first.company_name'),

        'base_salary_info': {
            'id': base_salary.id if base_salary else None,
            'annual_salary': base_salary.annual_salary if base_salary else None,
            'total_salary': base_salary.total_salary if base_salary else None,
            'total_deduction': base_salary.total_deduction if base_salary else None,
            'real_salary': base_salary.real_salary if base_salary else None,

            'base_salary': base_salary.base_salary if base_salary else None,
            'extend_salary': base_salary.extend_salary if base_salary else None,
            'night_salary': base_salary.night_salary if base_salary else None,
            'holiday_salary': base_salary.holiday_salary if base_salary else None,
            'meal_salary': base_salary.meal_salary if base_salary else None,
            'drive_salary': base_salary.drive_salary if base_salary else None,
            'bonus': base_salary.bonus if base_salary else None,
            'incentive': base_salary.incentive if base_salary else None,
            'etc_p1v': base_salary.etc_p1v if base_salary else None,
            'etc_p2v': base_salary.etc_p2v if base_salary else None,
            'etc_p3v': base_salary.etc_p3v if base_salary else None,
            'etc_p4v': base_salary.etc_p4v if base_salary else None,
            'etc_p5v': base_salary.etc_p5v if base_salary else None,

            'pension': base_salary.pension if base_salary else None,
            'health': base_salary.health if base_salary else None,
            'care': base_salary.care if base_salary else None,
            'unemployment': base_salary.unemployment if base_salary else None,
            'income_tax': base_salary.income_tax if base_salary else None,
            'local_tax': base_salary.local_tax if base_salary else None,
            'etc_m1v': base_salary.etc_m1v if base_salary else None,
            'etc_m2v': base_salary.etc_m2v if base_salary else None,
            'etc_m3v': base_salary.etc_m3v if base_salary else None,
            'etc_m4v': base_salary.etc_m4v if base_salary else None,
            'etc_m5v': base_salary.etc_m5v if base_salary else None,
            'etc_m6v': base_salary.etc_m6v if base_salary else None,

            'created_at': base_salary.created_at.strftime(
                '%Y-%m-%d %H:%M') if base_salary and base_salary.created_at else '',
            'updated_at': base_salary.updated_at.strftime(
                '%Y-%m-%d %H:%M') if base_salary and base_salary.updated_at else '',
            'created_by_id': safe_get(base_salary, 'created_by.id'),
            'created_by_name': safe_get(base_salary, 'created_by.name'),
            'updated_by_id': safe_get(base_salary, 'updated_by.id'),
            'updated_by_name': safe_get(base_salary, 'updated_by.name'),
            'updated_by_job_level': safe_get(base_salary, 'updated_by.job_level.name'),
            'company_id': safe_get(base_salary, 'company.id'),
            'company_name': safe_get(base_salary, 'company.company_info.first.company_name'),
        }
    }


def safe_get(obj, path, default=''):
    try:
        for attr in path.split('.'):
            obj = getattr(obj, attr)
            if obj is None:
                return default
        return obj
    except AttributeError:
        return default


def get_UBS_obj(obj):
    return {
        'id': obj.id,
        'user': get_user_info(obj.user) if obj.user else '',
        'annual_salary': obj.annual_salary,
        'total_salary': obj.total_salary,
        'total_deduction': obj.total_deduction,
        'real_salary': obj.real_salary,

        'base_salary': obj.base_salary,
        'extend_salary': obj.extend_salary,
        'night_salary': obj.night_salary,
        'holiday_salary': obj.holiday_salary,
        'meal_salary': obj.meal_salary,
        'drive_salary': obj.drive_salary,
        'bonus': obj.bonus,
        'incentive': obj.incentive,
        'etc_p1v': obj.etc_p1v,
        'etc_p2v': obj.etc_p2v,
        'etc_p3v': obj.etc_p3v,
        'etc_p4v': obj.etc_p4v,
        'etc_p5v': obj.etc_p5v,

        'pension': obj.pension,
        'health': obj.health,
        'care': obj.care,
        'unemployment': obj.unemployment,
        'income_tax': obj.income_tax,
        'local_tax': obj.local_tax,
        'etc_m1v': obj.etc_m1v,
        'etc_m2v': obj.etc_m2v,
        'etc_m3v': obj.etc_m3v,
        'etc_m4v': obj.etc_m4v,
        'etc_m5v': obj.etc_m5v,
        'etc_m6v': obj.etc_m6v,

        'created_at': obj.created_at.strftime('%Y-%m-%d %H:%M') if obj.created_at else '',
        'updated_at': obj.updated_at.strftime('%Y-%m-%d %H:%M') if obj.updated_at else '',
        'created_by_id': safe_get(obj, 'created_by.id', None),
        'created_by_name': safe_get(obj, 'created_by.name', ''),
        'updated_by_id': safe_get(obj, 'updated_by.id', None),
        'updated_by_name': safe_get(obj, 'updated_by.name', ''),
        'updated_by_job_level': safe_get(obj, 'updated_by.job_level.name', ''),
        'company_id': safe_get(obj, 'company.id', None),
        'company_name': safe_get(obj, 'company.company_info.first.company_name', ''),
    }


def get_user_info(user):
    return {
        'id': user.id,
        'name': user.name,
        'profile_image': user.profile_image.url if user.profile_image and user.profile_image.name else '',
        'team': user.team.name if user.team else '',
        'job_level': user.job_level.name if user.job_level else '',
    }


class UserSalary_ExcelExport(View):
    def get(self, request, *args, **kwargs):
        qs = UserBaseSalary.objects.filter(company=request.user.company).select_related('user')

        # 회사 커스텀 헤더명 (없으면 기본값 사용)
        company_info = request.user.company.company_info.first()
        etc_p1t = company_info.etc_p1t or '추가지급1'
        etc_p2t = company_info.etc_p2t or '추가지급2'
        etc_p3t = company_info.etc_p3t or '추가지급3'
        etc_p4t = company_info.etc_p4t or '추가지급4'

        etc_m1t = company_info.etc_m1t or '추가공제1'
        etc_m2t = company_info.etc_m2t or '추가공제2'
        etc_m3t = company_info.etc_m3t or '추가공제3'
        etc_m4t = company_info.etc_m4t or '추가공제4'
        etc_m5t = company_info.etc_m5t or '추가공제5'
        etc_m6t = company_info.etc_m6t or '추가공제6'

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '기본급여설정'

        # 헤더 정의
        headers = [
            '사번', '이름', '계약연봉', '월 기본급', '연장근로', '야간근로', '휴일근로', '식대', '운전보조', '상여금', '성과급',
            etc_p1t, etc_p2t, etc_p3t, etc_p4t, '국민연금', '건강보험', '장기요양', '고용보험', '소득세', '지방소득세',
            etc_m1t, etc_m2t, etc_m3t, etc_m4t, etc_m5t, etc_m6t
        ]

        # 스타일 정의
        header_font = Font(bold=True, color="FFFFFF")  # 흰색, 굵게
        header_fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")

        # 헤더 행 작성
        ws.append(headers)

        for obj in qs:
            ws.append([
                obj.user.employee_code if obj.user else '',
                obj.user.name if obj.user else '',
                obj.annual_salary,
                obj.base_salary,
                obj.extend_salary,
                obj.night_salary,
                obj.holiday_salary,
                obj.meal_salary,
                obj.drive_salary,
                obj.bonus,
                obj.incentive,
                obj.etc_p1v,
                obj.etc_p2v,
                obj.etc_p3v,
                obj.etc_p4v,

                obj.pension,
                obj.health,
                obj.care,
                obj.unemployment,
                obj.income_tax,
                obj.local_tax,
                obj.etc_m1v,
                obj.etc_m2v,
                obj.etc_m3v,
                obj.etc_m4v,
                obj.etc_m5v,
                obj.etc_m6v,
            ])

        # 헤더 스타일 적용
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="198754", end_color="198754", fill_type="solid")

        # 필터 적용
        ws.auto_filter.ref = ws.dimensions

        # 열 너비 자동 조정 (한글 보정 포함)
        for column_cells in ws.columns:
            max_length = 0
            col = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    cell_value = str(cell.value) if cell.value else ''
                    display_length = 0
                    for ch in cell_value:
                        # 한글이나 전각 문자는 길이 2로 간주
                        display_length += 2 if ord(ch) > 127 else 1
                    max_length = max(max_length, display_length)
                except:
                    pass
            adjusted_width = max_length + 3  # 여유 공간 포함
            ws.column_dimensions[col].width = adjusted_width

        # 회계 형식 적용 (숫자 열 전체에)
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=ws.max_column):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0'

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(buffer.read(),
                                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename=user_salary_export.xlsx'
        return response


@method_decorator(csrf_exempt, name='dispatch')
class UserSalary_ExcelImport(View):
    @transaction.atomic
    def post(self, request, *args, **kwargs):
        file = request.FILES.get('file')
        if not file:
            return JsonResponse({'error': True, 'message': '파일이 없습니다.'}, status=400)

        try:
            wb = openpyxl.load_workbook(file)
            ws = wb.active

            for idx, row in enumerate(ws.iter_rows(min_row=2), start=2):  # 헤더 제외
                emp_code = row[0].value
                if not emp_code:
                    continue

                try:
                    user = UserMaster.objects.get(employee_code=emp_code, company=request.user.company)
                except UserMaster.DoesNotExist:
                    continue

                # 엑셀 값 파싱
                annual_salary = row[2].value
                base_salary = row[3].value or 0
                extend_salary = row[4].value or 0
                night_salary = row[5].value or 0
                holiday_salary = row[6].value or 0
                meal_salary = row[7].value or 0
                drive_salary = row[8].value or 0
                bonus = row[9].value or 0
                incentive = row[10].value or 0

                etc_p1v = row[11].value or 0
                etc_p2v = row[12].value or 0
                etc_p3v = row[13].value or 0
                etc_p4v = row[14].value or 0

                pension = row[15].value or 0
                health = row[16].value or 0
                care = row[17].value or 0
                unemployment = row[18].value or 0
                income_tax = row[19].value or 0
                local_tax = row[20].value or 0

                etc_m1v = row[21].value or 0
                etc_m2v = row[22].value or 0
                etc_m3v = row[23].value or 0
                etc_m4v = row[24].value or 0
                etc_m5v = row[25].value or 0
                etc_m6v = row[26].value or 0

                total_salary = (
                    base_salary + extend_salary + night_salary + holiday_salary + meal_salary + drive_salary + bonus +
                    incentive + etc_p1v + etc_p2v + etc_p3v + etc_p4v
                )

                total_deduction = (
                    pension + health + care + unemployment + income_tax + local_tax +
                    etc_m1v + etc_m2v + etc_m3v + etc_m4v + etc_m5v + etc_m6v
                )

                real_salary = total_salary - total_deduction

                UserBaseSalary.objects.update_or_create(
                    user=user,
                    company=request.user.company,
                    defaults={
                        'annual_salary': annual_salary,
                        'base_salary': base_salary,
                        'extend_salary': extend_salary,
                        'night_salary': night_salary,
                        'holiday_salary': holiday_salary,
                        'meal_salary': meal_salary,
                        'drive_salary': drive_salary,
                        'bonus': bonus,
                        'incentive': incentive,
                        'etc_p1v': etc_p1v,
                        'etc_p2v': etc_p2v,
                        'etc_p3v': etc_p3v,
                        'etc_p4v': etc_p4v,

                        'pension': pension,
                        'health': health,
                        'care': care,
                        'unemployment': unemployment,
                        'income_tax': income_tax,
                        'local_tax': local_tax,
                        'etc_m1v': etc_m1v,
                        'etc_m2v': etc_m2v,
                        'etc_m3v': etc_m3v,
                        'etc_m4v': etc_m4v,
                        'etc_m5v': etc_m5v,
                        'etc_m6v': etc_m6v,

                        'total_salary': total_salary,
                        'total_deduction': total_deduction,
                        'real_salary': real_salary,

                        'updated_by': request.user,
                        'created_by': request.user,
                    }
                )

            return JsonResponse({'success': True, 'message': '엑셀 업로드 완료'})

        except Exception as e:
            transaction.set_rollback(True)
            return JsonResponse({'error': True, 'message': str(e)}, status=500)


class SalaryLabel_Update(View):
    def post(self, request, *args, **kwargs):
        field = request.POST.get("field", "").strip()
        new_text = request.POST.get("newText", "").strip()

        allowed_fields = [
            "etc_p1t", "etc_p2t", "etc_p3t", "etc_p4t",
            "etc_m1t", "etc_m2t", "etc_m3t", "etc_m4t", "etc_m5t", "etc_m6t"
        ]

        if field not in allowed_fields:
            return JsonResponse({"success": False, "error": "허용되지 않은 필드입니다."}, status=400)

        company_info = request.user.company.company_info.first()
        if not company_info:
            return JsonResponse({"success": False, "error": "회사 정보가 없습니다."}, status=404)

        setattr(company_info, field, new_text)
        company_info.save()

        return JsonResponse({"success": True})


# 급여대장 관리 함수 ------------------------------------------------------------------------------------------------------

class SalaryMaster_List(View):
    def get(self, request, *args, **kwargs):
        year_sch = request.GET.get("year_sch", "").strip()
        month_sch = request.GET.get("month_sch", "").strip()

        if not request.user.is_master:  # 관리자외 접근금지
            qs = None
        else:
            qs = SalaryMaster.objects.filter(company=request.user.company).order_by('year', 'month')

        if year_sch:
            qs = qs.filter(year=year_sch)

        if month_sch:
            qs = qs.filter(month=month_sch)

        # Pagination
        _page = int(request.GET.get('page', 1)) if request.GET.get('page', '1').isdigit() else 1
        _size = int(request.GET.get('page_size', 10)) if request.GET.get('page_size', '10').isdigit() else 10
        qs_ps = Pagenation(qs, _size, _page)

        pre_page = _page - 1
        url_pre = f"/?page_size={_size}&page={pre_page}" if pre_page >= 1 else None
        next_page = _page + 1
        url_next = f"/?page_size={_size}&page={next_page}" if next_page <= qs_ps.paginator.num_pages else None

        results = [get_SM_obj(row) for row in qs_ps]
        context = {
            'count': qs_ps.paginator.count,
            'previous': url_pre,
            'next': url_next,
            'results': results,
        }
        return JsonResponse(context, safe=False)


class SalaryMaster_Create(View):
    @transaction.atomic
    def post(self, request, *args, **kwargs):
        try:
            year = int(request.POST.get("year"))
            month = int(request.POST.get("month"))
            user_ids = request.POST.getlist("user_ids[]")

            try:
                salary_master = SalaryMaster.objects.create(
                    year=year,
                    month=month,
                    company=request.user.company,
                    created_by=request.user,
                    updated_by=request.user,
                )
            except IntegrityError:
                return JsonResponse({"result": "error", "message": f"{year}년 {month}월의 급여대장이 이미 존재합니다."}, status=400)

            for user_id in user_ids:
                try:
                    user = UserMaster.objects.get(id=user_id)
                    user_base = UserBaseSalary.objects.get(user=user)
                except (UserMaster.DoesNotExist, UserBaseSalary.DoesNotExist):
                    continue

                SalaryItem.objects.create(
                    salary=salary_master,
                    user=user,

                    base_salary=user_base.base_salary,
                    extend_salary=user_base.extend_salary,
                    night_salary=user_base.night_salary,
                    holiday_salary=user_base.holiday_salary,
                    meal_salary=user_base.meal_salary,
                    drive_salary=user_base.drive_salary,
                    bonus=user_base.bonus,
                    incentive=user_base.incentive,
                    etc_p1v=user_base.etc_p1v,
                    etc_p2v=user_base.etc_p2v,
                    etc_p3v=user_base.etc_p3v,
                    etc_p4v=user_base.etc_p4v,

                    pension=user_base.pension,
                    health=user_base.health,
                    care=user_base.care,
                    unemployment=user_base.unemployment,
                    income_tax=user_base.income_tax,
                    local_tax=user_base.local_tax,
                    etc_m1v=user_base.etc_m1v,
                    etc_m2v=user_base.etc_m2v,
                    etc_m3v=user_base.etc_m3v,
                    etc_m4v=user_base.etc_m4v,
                    etc_m5v=user_base.etc_m5v,
                    etc_m6v=user_base.etc_m6v,

                    total_salary=user_base.total_salary,
                    total_deduction=user_base.total_deduction,
                    real_salary=user_base.real_salary,

                    created_by=request.user,
                    updated_by=request.user,
                    company=request.user.company,
                )

                # 알림센터 메세지 전송
                noti_create_fn(
                    content=f"{salary_master.year}년 {salary_master.month}월 급여명세서",
                    user=user,
                    url=f"/my_salary/",
                    noti_type="salary_master_created"
                )

            return JsonResponse({"result": "ok", "msg": "급여대장이 생성되었습니다."})

        except Exception as e:
            return JsonResponse({"result": "error", "msg": str(e)}, status=500)


class SalaryMaster_Update(View):
    @transaction.atomic
    def post(self, request, *args, **kwargs):
        try:
            pk = request.POST.get('pk', '')
            desc1 = request.POST.get('desc1', '')
            desc2 = request.POST.get('desc2', '')
            desc3 = request.POST.get('desc3', '')

            obj = get_object_or_404(SalaryMaster, pk=int(pk))
            obj.desc1 = desc1
            obj.desc2 = desc2
            obj.desc3 = desc3

            obj.updated_by = request.user
            obj.updated_at = timezone.now()
            obj.save()

            return JsonResponse({"result": "ok", "msg": "급여대장의 추가정보가 수정되었습니다."})

        except Exception as e:
            transaction.set_rollback(True)
            msg = get_excep_msg(e)
            return JsonResponse({'error': True, 'message': msg}, status=500)


class SalaryMaster_Delete(View):
    @transaction.atomic
    def post(self, request):
        try:
            pk = request.POST.get('pk', '')
            if (pk == ''):
                return JsonResponse({'error': True, 'message': "삭제할 항목이 선택되지 않았습니다."})

            obj = get_object_or_404(SalaryMaster, pk=int(pk))
            obj.delete()

            context = {'success': True}
            return JsonResponse(context, safe=False)

        except Exception as e:
            transaction.set_rollback(True)
            msg = get_excep_msg(e)
            return JsonResponse({'error': True, 'message': msg}, status=500)


def get_SM_obj(obj):
    items = obj.items.select_related('user', 'user__job_title', 'user__job_level', 'user__team').all()

    total_user = 0
    total_payment = 0
    total_deduction = 0
    total_real_salary = 0

    item_list = []
    for item in items:
        user = item.user
        total_user += 1

        # 지급 항목 합산
        payment = sum(filter(None, [
            item.base_salary, item.extend_salary, item.night_salary, item.holiday_salary,
            item.meal_salary, item.drive_salary, item.bonus, item.incentive,
            item.etc_p1v, item.etc_p2v, item.etc_p3v, item.etc_p4v,
        ]))

        # 공제 항목 합산
        deduction = sum(filter(None, [
            item.pension, item.health, item.care, item.unemployment,
            item.income_tax, item.local_tax,
            item.etc_m1v, item.etc_m2v, item.etc_m3v,
            item.etc_m4v, item.etc_m5v, item.etc_m6v,
        ]))

        real_salary = item.real_salary or 0

        total_payment += payment
        total_deduction += deduction
        total_real_salary += real_salary

        try:
            base_info = user.base_salary.latest('created_at')
        except UserBaseSalary.DoesNotExist:
            base_info = None

        item_list.append({
            'user_id': user.id,
            'name': user.name,
            'employee_code': user.employee_code,
            'job_title_name': user.job_title.name if user.job_title else '',
            'job_level_name': user.job_level.name if user.job_level else '',
            'team_name': user.team.name if user.team else '',
            'profile_image': user.profile_image.url if user.profile_image and user.profile_image.name else '',

            'base_salary': item.base_salary,
            'extend_salary': item.extend_salary,
            'night_salary': item.night_salary,
            'holiday_salary': item.holiday_salary,
            'meal_salary': item.meal_salary,
            'drive_salary': item.drive_salary,
            'bonus': item.bonus,
            'incentive': item.incentive,
            'pension': item.pension,
            'health': item.health,
            'care': item.care,
            'unemployment': item.unemployment,
            'income_tax': item.income_tax,
            'local_tax': item.local_tax,
            'total_deduction': item.total_deduction,
            'real_salary': item.real_salary,

            'etc_p1v': item.etc_p1v,
            'etc_p2v': item.etc_p2v,
            'etc_p3v': item.etc_p3v,
            'etc_p4v': item.etc_p4v,
            'etc_p5v': item.etc_p5v,
            'etc_m1v': item.etc_m1v,
            'etc_m2v': item.etc_m2v,
            'etc_m3v': item.etc_m3v,
            'etc_m4v': item.etc_m4v,
            'etc_m5v': item.etc_m5v,
            'etc_m6v': item.etc_m6v,

            'base_salary_info': {
                'annual_salary': base_info.annual_salary if base_info else None,
                'base_salary': base_info.base_salary if base_info else None,
                'updated_at': base_info.updated_at.strftime(
                    '%Y-%m-%d %H:%M') if base_info and base_info.updated_at else ''
            }
        })

    return {
        'salary_id': obj.id,
        'year': obj.year,
        'month': obj.month,
        'desc1': obj.desc1,
        'desc2': obj.desc2,
        'desc3': obj.desc3,
        'created_at': obj.created_at.strftime('%Y-%m-%d %H:%M') if obj.created_at else '',
        'created_by': obj.created_by.name if obj.created_by else '',
        'total_user': total_user,
        'total_payment': total_payment,
        'total_deduction': total_deduction,
        'total_real_salary': total_real_salary,
        'items': item_list,
    }


# 개인 급여명세서 호출 함수 ------------------------------------------------------------------------------------------------

@login_required
@csrf_exempt
def my_salary_check_fn(request):
    if request.method == 'POST':
        data = json.loads(request.body)
        password = data.get('password')
        if request.user.check_password(password):
            return JsonResponse({'ok': True})
        return JsonResponse({'ok': False}, status=403)


class MySalary_List(View):
    def get(self, request, *args, **kwargs):
        user = request.GET.get("user", "").strip()
        data = request.GET.get("data", "").strip()

        if not str(request.user.id) == user:
            qs = []
        else:
            qs = SalaryItem.objects.filter(user=request.user, id=data).order_by('id')

        results = [get_MY_obj(row) for row in qs]
        context = {
            'results': results,
        }
        return JsonResponse(context, safe=False)


def get_MY_obj(obj):
    return {
        'id': obj.salary_id,
        'user': get_user_info(obj.user) if obj.user else '',
        'salary_year': obj.salary.year,
        'salary_month': obj.salary.month,
        'salary_desc1': obj.salary.desc1,
        'salary_desc2': obj.salary.desc2,
        'real_salary': obj.real_salary,

        'base_salary': obj.base_salary,
        'extend_salary': obj.extend_salary,
        'night_salary': obj.night_salary,
        'holiday_salary': obj.holiday_salary,
        'meal_salary': obj.meal_salary,
        'drive_salary': obj.drive_salary,
        'bonus': obj.bonus,
        'incentive': obj.incentive,
        'etc_p1v': obj.etc_p1v,
        'etc_p2v': obj.etc_p2v,
        'etc_p3v': obj.etc_p3v,
        'etc_p4v': obj.etc_p4v,
        'etc_p5v': obj.etc_p5v,
        'total_salary': sum(getattr(obj, field) or 0 for field in [
            'base_salary', 'extend_salary', 'night_salary', 'holiday_salary', 'meal_salary',
            'drive_salary', 'bonus', 'incentive', 'etc_p1v', 'etc_p2v', 'etc_p3v', 'etc_p4v', 'etc_p5v'
        ]),

        'pension': obj.pension,
        'health': obj.health,
        'care': obj.care,
        'unemployment': obj.unemployment,
        'income_tax': obj.income_tax,
        'local_tax': obj.local_tax,
        'etc_m1v': obj.etc_m1v,
        'etc_m2v': obj.etc_m2v,
        'etc_m3v': obj.etc_m3v,
        'etc_m4v': obj.etc_m4v,
        'etc_m5v': obj.etc_m5v,
        'etc_m6v': obj.etc_m6v,
        'total_deduction': sum(getattr(obj, field) or 0 for field in [
            'pension', 'health', 'care', 'unemployment', 'income_tax', 'local_tax',
            'etc_m1v', 'etc_m2v', 'etc_m3v', 'etc_m4v', 'etc_m5v', 'etc_m6v'
        ]),

        'created_at': obj.created_at.strftime('%Y-%m-%d %H:%M') if obj.created_at else '',
        'updated_at': obj.updated_at.strftime('%Y-%m-%d %H:%M') if obj.updated_at else '',
        'created_by_id': safe_get(obj, 'created_by.id', None),
        'created_by_name': safe_get(obj, 'created_by.name', ''),
        'updated_by_id': safe_get(obj, 'updated_by.id', None),
        'updated_by_name': safe_get(obj, 'updated_by.name', ''),
        'updated_by_job_level': safe_get(obj, 'updated_by.job_level.name', ''),
        'company_id': safe_get(obj, 'company.id', None),
        'company_name': safe_get(obj, 'company.company_info.first.company_name', ''),
    }
